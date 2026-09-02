#!/usr/bin/env python
# coding: utf-8

"""
Step 9: Analyze daily money flow data.
Source Code: Daily_TA4C_Analyze_MoneyFlow_mproc.py
Input data:     hk_daily_moneyflow_analysis
Output data:    daily_moneyflow_analysis.pdf
Function:
1. 生成资金流可视化图表，包括机构/散户资金净流入/净流出。
2. 列举关键资金流特征指标。
Modified: 从 SQLite 数据库读取数据 - 使用小写列名
"""

# === 多进程优化导入 ===
import multiprocessing as mp
from multiprocessing import Pool, cpu_count
import concurrent.futures
from functools import partial

# === 设置 Matplotlib 为非交互式后端，避免多进程问题 ===
import matplotlib
matplotlib.use('Agg')  # 使用非GUI后端
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import matplotlib.gridspec as gridspec
from matplotlib.backends.backend_pdf import PdfPages
import seaborn as sns

# 核心库
import os
import sys
import io
import time
import warnings
import logging
import traceback
import datetime
from datetime import datetime, timedelta
import sqlite3

# 数据处理
import pandas as pd
import numpy as np

# 专业金融图表
import mplfinance as mpf

# PDF处理
from fpdf import FPDF
from PyPDF2 import PdfReader, PdfWriter

# Excel处理
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# 其他
from pathlib import Path
from scipy import stats

# === 在导入语句之后，添加警告过滤 ===
import warnings

# 抑制 fpdf 字体相关的警告
warnings.filterwarnings("ignore", category=UserWarning, module='fpdf.ttfonts')
# 抑制其他 UserWarning（可选）
warnings.filterwarnings("ignore", category=UserWarning)


# === 添加UTL路径到系统路径 ===
# 当前脚本所在目录：
script_dir = Path(__file__).resolve().parent

# 将 Core 目录加入 sys.path，以便导入 utl 包
core_dir = script_dir.parent   # TA_Workflow/Core
if str(core_dir) not in sys.path:
    sys.path.insert(0, str(core_dir))

# 从 utl.stock_analysis_utl 导入 get_project_root（该函数返回 Core 目录）
from utl.stock_analysis_utl import get_project_root

# 利用该函数获取 Core 目录，再取其父目录得到项目根目录 (TA_Workflow)
core_dir_from_utils = get_project_root()   # Path 对象，指向 Core
project_dir = core_dir_from_utils.parent  # TA_Workflow

# 确保项目根目录也在 sys.path（便于其他模块使用，虽非必需）
if str(project_dir) not in sys.path:
    sys.path.insert(0, str(project_dir))

# 从 utl.stock_analysis_utl 导入所需工具（同在 utl 目录下）
try:
    from utl.stock_analysis_utl import (
        load_config,
        setup_windows_encoding,
        GlobalConfig,
        get_validated_dates
    )
except ImportError as e:
    print("=" * 60)
    print("❌ 严重错误：无法导入 utl.stock_analysis_utl 模块")
    print(f"   期望路径: {core_dir / 'utl' / 'stock_analysis_utl.py'}")
    print(f"   请确认该文件存在，且 utl 包中有 __init__.py")
    print(f"   详细错误: {e}")
    print("=" * 60)
    sys.exit(1)
    

# === 强制UTF-8编码输出 ===
if 'get_ipython' not in globals():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')
else:
    # 在Jupyter环境中添加兼容处理
    import ipykernel
    if ipykernel:
        sys.stdout.encoding = 'utf-8'
        sys.stderr.encoding = 'utf-8'

# === 设置日志（在导入后立即初始化）===
logger = logging.getLogger(__name__)

# 执行Windows编码设置
try:
    setup_windows_encoding()
except NameError:
    logger.warning("setup_windows_encoding函数未定义，跳过编码设置")

# === 多进程优化配置 ===
class ProcessConfig:
    """多进程配置类"""
    MAX_WORKERS = max(1, cpu_count() - 1)  # 保留一个CPU核心
    CHUNK_SIZE = 5  # 每个进程处理的股票数量
    TIMEOUT = 300  # 进程超时时间（秒）

# === 数据库工具函数 ===
def get_db_connection(db_path=None):
    """获取数据库连接"""
    if db_path is None:
        # 默认数据库路径
        db_path = os.path.join(project_dir, 'SQLiteDB', 'HK_Stock.db')
    
    if not os.path.exists(db_path):
        logger.error(f"数据库文件不存在: {db_path}")
        return None
    
    try:
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        return conn
    except Exception as e:
        logger.error(f"连接数据库失败: {e}")
        return None

def safe_divide(numerator, denominator, default=0):
    """安全除法，避免除零错误"""
    try:
        if denominator == 0 or pd.isna(denominator):
            return default
        return numerator / denominator
    except:
        return default

def get_stock_analysis_data(ticker, start_date=None, end_date=None, db_path=None):
    """
    从数据库获取股票分析数据 - 直接使用小写列名，不进行映射
    """
    conn = get_db_connection(db_path)
    if conn is None:
        return None
    
    try:
        # 构建SQL查询 - 使用小写列名
        query = """
        SELECT 
            k.date,
            k.stock_code,
            k.stock_name,
            k.open,
            k.high,
            k.low,
            k.close,
            k.volume,
            k.amount,
            k.turnover_rate,
            k.change_percent,
            k.change_amount,
            k.avg_volume_5d,
            k.volume_ratio5,
            k.avg_volume_20d,
            k.volume_ratio20,
            k.turnover_rate_75th_percentile_500d,
            k.ema5,
            k.ema10,
            k.ema20,
            k.ema50,
            k.ema60,
            k.ema120,
            k.ema200,
            k.ema250,
            k.bias5,
            k.bias10,
            k.bias20,
            k.bias50,
            k.bias60,
            k.bias120,
            k.bias200,
            k.bias250,
            k.volume_ratio5_min_500d,
            k.volume_ratio5_max_500d,
            k.volume_ratio5_10th_percentile_500d,
            k.volume_ratio5_25th_percentile_500d,
            k.volume_ratio5_50th_percentile_500d,
            k.volume_ratio5_75th_percentile_500d,
            k.volume_ratio5_90th_percentile_500d,
            k.volume_ratio20_min_500d,
            k.volume_ratio20_max_500d,
            k.volume_ratio20_10th_percentile_500d,
            k.volume_ratio20_25th_percentile_500d,
            k.volume_ratio20_50th_percentile_500d,
            k.volume_ratio20_75th_percentile_500d,
            k.volume_ratio20_90th_percentile_500d,
            m.extra_large,
            m.large,
            m.medium,
            m.small,
            m.institutional_flow,
            m.individual_flow,
            m.inst_5ma,
            m.ind_5ma,
            m.idr,
            m.fbi,
            m.capital_trend,
            m.price_chgpct
        FROM hk_daily_kline_analysis k
        INNER JOIN hk_daily_moneyflow_analysis m
            ON k.stock_code = m.stock_code AND k.date = m.date
        WHERE k.stock_code = ?
        """
        
        params = [ticker]
        
        # 添加日期过滤
        if start_date:
            query += " AND k.date >= ?"
            params.append(start_date)
        if end_date:
            query += " AND k.date <= ?"
            params.append(end_date)
        
        query += " ORDER BY k.date ASC"
        
        # 执行查询 - 直接使用小写列名
        df = pd.read_sql_query(query, conn, params=params)
        
        if df.empty:
            logger.warning(f"股票 {ticker} 在指定日期范围内没有数据")
            return None
        
        # 确保日期列为datetime类型
        df['date'] = pd.to_datetime(df['date'])
        
        # 验证必要的列是否存在（使用小写）
        required_cols = ['institutional_flow', 'individual_flow']
        missing_cols = [col for col in required_cols if col not in df.columns]
        if missing_cols:
            logger.error(f"缺少必要的资金流列: {missing_cols}")
            logger.error(f"当前列名: {df.columns.tolist()}")
            return None
        
        # 确保资金流列没有空值 - 用0填充
        flow_cols = ['institutional_flow', 'individual_flow', 'extra_large', 'large', 'medium', 'small']
        for col in flow_cols:
            if col in df.columns:
                df[col] = df[col].fillna(0)
        
        # 计算 institutional_net_ratio 和 individual_net_ratio（小写）
        if 'idr' in df.columns:
            df['institutional_net_ratio'] = df['idr'] * 100
            df['individual_net_ratio'] = (1 - df['idr']) * 100
        else:
            # 如果没有 IDR，使用资金流数据计算
            if 'institutional_flow' in df.columns and 'individual_flow' in df.columns:
                total_flow = df['institutional_flow'] + df['individual_flow']
                df['institutional_net_ratio'] = np.where(
                    total_flow != 0,
                    df['institutional_flow'] / total_flow * 100,
                    50
                )
                df['individual_net_ratio'] = 100 - df['institutional_net_ratio']
            else:
                df['institutional_net_ratio'] = 50
                df['individual_net_ratio'] = 50
        
        # 计算各种比率 - 使用小写列名和安全除法
        if 'extra_large' in df.columns and 'institutional_flow' in df.columns:
            df['extra_large_ratio'] = df.apply(
                lambda row: safe_divide(row['extra_large'], row['institutional_flow'], 0),
                axis=1
            )
        else:
            df['extra_large_ratio'] = 0
            
        if 'large' in df.columns and 'institutional_flow' in df.columns:
            df['large_ratio'] = df.apply(
                lambda row: safe_divide(row['large'], row['institutional_flow'], 0),
                axis=1
            )
        else:
            df['large_ratio'] = 0
            
        if 'medium' in df.columns and 'individual_flow' in df.columns:
            df['medium_ratio'] = df.apply(
                lambda row: safe_divide(row['medium'], row['individual_flow'], 0),
                axis=1
            )
        else:
            df['medium_ratio'] = 0
            
        if 'small' in df.columns and 'individual_flow' in df.columns:
            df['small_ratio'] = df.apply(
                lambda row: safe_divide(row['small'], row['individual_flow'], 0),
                axis=1
            )
        else:
            df['small_ratio'] = 0
        
        # 计算移动平均（小写）
        if 'inst_5ma' not in df.columns and 'institutional_flow' in df.columns:
            df['inst_5ma'] = df['institutional_flow'].rolling(window=5, min_periods=1).mean()
        elif 'inst_5ma' not in df.columns:
            df['inst_5ma'] = 0
            
        if 'ind_5ma' not in df.columns and 'individual_flow' in df.columns:
            df['ind_5ma'] = df['individual_flow'].rolling(window=5, min_periods=1).mean()
        elif 'ind_5ma' not in df.columns:
            df['ind_5ma'] = 0
        
        # 确保 volume_ratio5 存在
        if 'volume_ratio5' not in df.columns:
            df['volume_ratio5'] = 1.0
        
        logger.info(f"成功获取股票 {ticker} 的数据，共 {len(df)} 条记录")
        logger.info(f"列名（小写）: {df.columns.tolist()}")
        return df
        
    except Exception as e:
        logger.error(f"获取股票 {ticker} 数据失败: {e}")
        import traceback
        traceback.print_exc()
        return None
    finally:
        conn.close()
        
def get_unit_and_scale(data_series):
    """根据数据量级返回合适的单位和缩放因子（Mil/Bil），增加错误处理"""
    try:
        # 检查数据是否有效
        if data_series is None or len(data_series) == 0:
            logger.warning("数据序列为空，使用默认单位 Mil")
            return 'Mil', 1e6

        # 过滤掉NaN和无穷大的值
        clean_series = data_series.replace([np.inf, -np.inf], np.nan).dropna()

        if len(clean_series) == 0:
            logger.warning("数据序列全部为无效值，使用默认单位 Mil")
            return 'Mil', 1e6

        # 计算最大值
        max_abs = max(abs(clean_series.min()), abs(clean_series.max()))

        if max_abs >= 1e9:
            return 'Bil', 1e9
        else:
            return 'Mil', 1e6
    except Exception as e:
        logger.error(f"计算单位和缩放因子时出错: {e}，使用默认值")
        return 'Mil', 1e6

# === 字体相关函数 ===
def find_font_file(font_names):
    """查找系统中可用的中文字体文件路径，优先返回TTF格式"""
    # 优先查找TTF字体
    ttf_fonts = [name for name in font_names if name.lower().endswith('.ttf')]
    if ttf_fonts:
        for font_name in ttf_fonts:
            try:
                if os.name == 'nt':  # Windows
                    font_path = os.path.join(os.environ['WINDIR'], 'Fonts', font_name)
                    if os.path.exists(font_path):
                        return font_path
                elif os.name == 'posix':  # macOS/Linux
                    search_paths = [
                        '/Library/Fonts/',
                        '/System/Library/Fonts/',
                        '/usr/share/fonts/',
                        os.path.expanduser('~/Library/Fonts/')
                    ]
                    for path in search_paths:
                        font_path = os.path.join(path, font_name)
                        if os.path.exists(font_path):
                            return font_path
            except Exception:
                continue

    # 尝试通过font_manager查找
    try:
        import matplotlib.font_manager as fm
        for font in fm.findSystemFonts():
            font_name = os.path.basename(font).lower()
            if any(name in font_name for name in ['simhei', 'msyh', 'simsun', 'stkai', 'simfang']):
                return font
    except Exception:
        pass

    return None

class ChinesePDF(FPDF):
    """支持中文字符的PDF生成类，修复所有编码问题"""
    def __init__(self, font_path=None, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.font_path = font_path
        self.font_name = "customfont"
        self.font_added = False

        if self.font_path and os.path.exists(self.font_path):
            self._add_custom_font()

    def _add_custom_font(self):
        """添加自定义字体到PDF"""
        try:
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", category=UserWarning)
                self.add_font(self.font_name, '', self.font_path, uni=True)
                self.add_font(self.font_name, 'B', self.font_path, uni=True)
                self.add_font(self.font_name, 'I', self.font_path, uni=True)
                self.add_font(self.font_name, 'BI', self.font_path, uni=True)
                self.font_added = True
                logger.info(f"已添加自定义字体: {self.font_path}")
        except Exception as e:
            logger.error(f"添加自定义字体失败: {e}")
            self.font_added = False

    def header(self):
        """自定义页眉"""
        if self.font_added:
            self.set_font(self.font_name, 'B', 15)
        else:
            self.set_font('Arial', 'B', 15)

        title = getattr(self, 'title', '资金流动分析报告')
        title_w = self.get_string_width(title) + 6
        self.set_x((210 - title_w) / 2)
        self.set_draw_color(0, 80, 180)
        self.set_fill_color(230, 230, 0)
        self.set_text_color(0, 0, 0)
        self.cell(title_w, 10, title, border=1, ln=1, align='C', fill=1)
        self.ln(10)

    def footer(self):
        """自定义页脚"""
        self.set_y(-15)
        if self.font_added:
            self.set_font(self.font_name, 'I', 8)
        else:
            self.set_font('Arial', 'I', 8)
        self.set_text_color(128)
        self.cell(0, 10, f'第 {self.page_no()} 页', 0, 0, 'C')

    def add_text(self, text, font_size=12, style=''):
        """添加支持中文的文本"""
        try:
            if self.font_added:
                self.set_font(self.font_name, style, font_size)
            else:
                if style == 'B':
                    self.set_font('Arial', 'B', font_size)
                elif style == 'I':
                    self.set_font('Arial', 'I', font_size)
                elif style == 'BI':
                    self.set_font('Arial', 'BI', font_size)
                else:
                    self.set_font('Arial', '', font_size)

            if not isinstance(text, str):
                text = str(text)

            self._safe_multi_cell(0, 5, text)
            self.ln(5)
            return True
        except Exception as e:
            logger.error(f"添加文本失败: {e}")
            return False

    def _safe_multi_cell(self, w, h, txt, border=0, align='J', fill=False):
        """安全的多行文本添加方法"""
        try:
            import unicodedata
            self.multi_cell(w, h, txt, border, align, fill)
        except UnicodeEncodeError:
            normalized_txt = unicodedata.normalize('NFC', txt)
            try:
                self.multi_cell(w, h, normalized_txt, border, align, fill)
            except Exception:
                ascii_txt = txt.encode('ascii', 'replace').decode('ascii')
                self.multi_cell(w, h, ascii_txt, border, align, fill)
        except Exception as e:
            logger.error(f"添加文本时出错: {e}")
            self.multi_cell(w, h, "[文本添加失败]", border, align, fill)

    def add_plot(self, image_path, width=180):
        """添加图片到PDF"""
        if os.path.exists(image_path):
            try:
                self.image(image_path, x=10, y=None, w=width)
                self.ln(5)
                return True
            except Exception as e:
                logger.error(f"添加图片失败: {e}")
                return False
        else:
            logger.warning(f"图片文件 '{image_path}' 不存在，跳过添加")
            return False

    def title_section(self, title):
        """添加标题部分"""
        if self.font_added:
            self.set_font(self.font_name, 'B', 16)
        else:
            self.set_font('Arial', 'B', 16)
        self.cell(0, 10, title, 0, 1, 'C')
        self.ln(5)

def close_pdf_reader(file_path):
    """尝试关闭PDF阅读器（Windows专用）"""
    if os.name != 'nt':
        return False

    try:
        import win32com.client
        wmi = win32com.client.GetObject("winmgmts:")
        processes = wmi.InstancesOf("Win32_Process")
        for process in processes:
            cmd = process.Properties_("CommandLine").Value or ""
            if f'"{file_path}"' in cmd or file_path in cmd:
                process.Terminate()
                return True
    except ImportError:
        logger.warning("win32com.client未安装，无法自动关闭PDF阅读器")
    except Exception as e:
        logger.error(f"关闭PDF阅读器时出错: {e}")

    return False

def open_pdf(pdf_path):
    """根据操作系统打开PDF文件"""
    if os.path.exists(pdf_path):
        try:
            if sys.platform == 'win32':
                os.startfile(pdf_path)
            elif sys.platform == 'darwin':
                import subprocess
                subprocess.call(('open', pdf_path))
            else:
                import subprocess
                subprocess.call(('xdg-open', pdf_path))
            return True
        except Exception as e:
            logger.error(f"打开PDF文件时出错: {str(e)}")
            return False
    else:
        logger.error(f"PDF文件 '{pdf_path}' 不存在")
        return False

# === 核心分析函数 ===
def create_grouped_stacked_bars(current_ticker, date_range_label, period_data_cache, data_periods=None):
    """
    Optimized function to create grouped stacked bar charts with negative value support
    """
    if data_periods is None:
        data_periods = [20, 60, 'AP']

    periods = []
    for period in data_periods:
        if period in period_data_cache:
            periods.append(period_data_cache[period])

    if not periods:
        return None

    investor_types = ['extra_large', 'large', 'medium', 'small']
    colors = ['#1f77b4', '#5ca0d3', '#ff7f0e', '#ffaa5c']

    all_values = []
    for _, period_df in periods:
        for itype in investor_types:
            if itype in period_df.columns:
                all_values.append(period_df[itype].sum())

    if not all_values:
        return None

    total_unit, total_scale = get_unit_and_scale(pd.Series(all_values))

    plot_data = []
    for period_name, period_df in periods:
        extra_large = period_df['extra_large'].sum() / total_scale if 'extra_large' in period_df.columns else 0
        large = period_df['large'].sum() / total_scale if 'large' in period_df.columns else 0
        medium = period_df['medium'].sum() / total_scale if 'medium' in period_df.columns else 0
        small = period_df['small'].sum() / total_scale if 'small' in period_df.columns else 0

        period_data = {
            'Period': period_name,
            'Institutional_Flow': {
                'extra_large': extra_large,
                'large': large,
            },
            'Individual_Flow': {
                'medium': medium,
                'small': small
            },
            'Total Institutional': extra_large + large,
            'Total Individual': medium + small
        }
        plot_data.append(period_data)

    fig, ax = plt.subplots(figsize=(16, 9))

    x = np.arange(len(plot_data))
    width = 0.35
    offset = width / 2

    min_y = 0
    max_y = 0
    for period in plot_data:
        inst_components = period['Institutional_Flow']
        inst_bottom = 0
        for itype in ['extra_large', 'large']:
            value = inst_components[itype]
            low = min(inst_bottom, inst_bottom + value)
            high = max(inst_bottom, inst_bottom + value)
            min_y = min(min_y, low)
            max_y = max(max_y, high)
            inst_bottom += value

        ind_components = period['Individual_Flow']
        ind_bottom = 0
        for itype in ['medium', 'small']:
            value = ind_components[itype]
            low = min(ind_bottom, ind_bottom + value)
            high = max(ind_bottom, ind_bottom + value)
            min_y = min(min_y, low)
            max_y = max(max_y, high)
            ind_bottom += value

    padding = 0.15 * (max_y - min_y) if max_y != min_y else 0.15
    ax.set_ylim(min_y - padding, max_y + padding)

    y_range = max_y - min_y
    if y_range == 0:
        label_offset = 0.01 * abs(max_y) if max_y != 0 else 0.01
    else:
        label_offset = y_range * 0.01

    for period_idx, period in enumerate(plot_data):
        inst_components = [
            ('extra_large', period['Institutional_Flow']['extra_large'], colors[0]),
            ('large', period['Institutional_Flow']['large'], colors[1])
        ]

        inst_bottom = 0
        for component, value, color in inst_components:
            if value != 0:
                bar = ax.bar(
                    x[period_idx] - offset, 
                    value, 
                    width, 
                    bottom=inst_bottom,
                    color=color,
                    edgecolor='white'
                )
                mid_point = inst_bottom + value/2
                ax.text(
                    x[period_idx] - offset, 
                    mid_point, 
                    f'{value:,.1f}',
                    ha='center', 
                    va='center',
                    color='white',
                    fontweight='bold',
                    fontsize=9
                )
                inst_bottom += value

        ind_components = [
            ('medium', period['Individual_Flow']['medium'], colors[2]),
            ('small', period['Individual_Flow']['small'], colors[3])
        ]

        ind_bottom = 0
        for component, value, color in ind_components:
            if value != 0:
                bar = ax.bar(
                    x[period_idx] + offset, 
                    value, 
                    width, 
                    bottom=ind_bottom,
                    color=color,
                    edgecolor='white'
                )
                mid_point = ind_bottom + value/2
                ax.text(
                    x[period_idx] + offset, 
                    mid_point, 
                    f'{value:,.1f}',
                    ha='center', 
                    va='center',
                    color='white',
                    fontweight='bold',
                    fontsize=9
                )
                ind_bottom += value

        total_inst = period['Total Institutional']
        if total_inst != 0:
            va = 'bottom' if total_inst >= 0 else 'top'
            y_pos = total_inst + label_offset if total_inst >= 0 else total_inst - label_offset
            ax.text(
                x[period_idx] - offset, 
                y_pos, 
                f'Total: {total_inst:,.1f}',
                ha='center', 
                va=va,
                fontsize=9,
                fontweight='bold',
                color='#1f77b4'
            )

        total_ind = period['Total Individual']
        if total_ind != 0:
            va = 'bottom' if total_ind >= 0 else 'top'
            y_pos = total_ind + label_offset if total_ind >= 0 else total_ind - label_offset
            ax.text(
                x[period_idx] + offset, 
                y_pos, 
                f'Total: {total_ind:,.1f}',
                ha='center', 
                va=va,
                fontsize=9,
                fontweight='bold',
                color='#ff7f0e'
            )

    ax.set_title(
        f'{current_ticker} Investor Flow Breakdown Comparison',
        fontsize=16,
        pad=20
    )
    ax.set_ylabel(f'Total Flow Amount ({total_unit})', fontsize=13)
    ax.set_xticks(x)
    ax.set_xticklabels([p['Period'] for p in plot_data], fontsize=11)
    ax.grid(axis='y', alpha=0.3)

    for i in range(1, len(plot_data)):
        ax.axvline(x=i - 0.5, color='gray', linestyle='-', alpha=0.5)

    from matplotlib.patches import Patch
    legend_elements = [
        Patch(facecolor=colors[0], label='Institutional: extra_large'),
        Patch(facecolor=colors[1], label='Institutional: large'),
        Patch(facecolor=colors[2], label='Individual: medium'),
        Patch(facecolor=colors[3], label='Individual: small'),
    ]

    ax.legend(
        handles=legend_elements, 
        loc='upper center',
        bbox_to_anchor=(0.5, 1.15),
        ncol=4,
        frameon=True,
        framealpha=1,
        fontsize=10
    )

    fig.text(0.5, 0.01, 'Data Source: Flow Analysis System', 
             ha='center', fontsize=9, alpha=0.7)

    plt.tight_layout()
    safe_periods = [str(p).replace(':', '-') for p in data_periods]
    period_str = "_".join(safe_periods)
    plot_filename = os.path.join(GlobalConfig.full_temp_dir, f"{current_ticker}_grouped_stacked_breakdown_{period_str}.png")
    plt.savefig(plot_filename, dpi=150, bbox_inches='tight')
    plt.close(fig)
    return plot_filename


def analyze_investor_dominance_from_db(ticker, pdf_report, start_date=None, end_date=None, db_path=None):
    """
    从数据库读取数据进行分析 - 使用小写列名
    """
    output_dir = os.path.dirname(pdf_report)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir, exist_ok=True)

    plot_filenames = []
    pdf_generated = False
    current_ticker = ticker

    try:
        df = get_stock_analysis_data(ticker, start_date, end_date, db_path)
        
        if df is None or df.empty:
            raise ValueError(f"无法获取股票 {ticker} 的数据")

        df["date"] = pd.to_datetime(df["date"], errors="coerce")
        
        required_cols = ['date', 'institutional_flow', 'individual_flow', 'extra_large', 'large', 'medium', 'small']
        
        missing_cols = [col for col in required_cols if col not in df.columns]
        if missing_cols:
            logger.error(f"缺少必要的列: {missing_cols}")
            logger.error(f"当前所有列: {df.columns.tolist()}")
            return False
        
        logger.info(f"原始数据行数: {len(df)}")
        
        for col in ['institutional_flow', 'individual_flow', 'extra_large', 'large', 'medium', 'small']:
            if col in df.columns:
                df[col] = df[col].fillna(0)
        
        df = df.dropna(subset=['date'])
        df = df.sort_values(by="date", ascending=True)

        if len(df) == 0:
            logger.warning("Dataframe is empty after preprocessing")
            return False

        logger.info(f"预处理后数据行数: {len(df)}")
        logger.info(f"日期范围: {df['date'].min()} 到 {df['date'].max()}")

        first_day = df['date'].min()
        last_day = df['date'].max()

        if start_date:
            start_date_dt = pd.to_datetime(start_date)
            if start_date_dt < first_day:
                logger.warning(f"Start date {start_date_dt.strftime('%Y-%m-%d')} is earlier than first data day {first_day.strftime('%Y-%m-%d')}, using first day")
                start_date_dt = first_day
        else:
            start_date_dt = first_day

        if end_date:
            end_date_dt = pd.to_datetime(end_date)
            if end_date_dt > last_day:
                logger.warning(f"End date {end_date_dt.strftime('%Y-%m-%d')} is later than last data day {last_day.strftime('%Y-%m-%d')}, using last day")
                end_date_dt = last_day
        else:
            end_date_dt = last_day

        date_range_label = f"{start_date_dt.strftime('%Y-%m-%d')} to {end_date_dt.strftime('%Y-%m-%d')}"

        mask = (df['date'] >= start_date_dt) & (df['date'] <= end_date_dt)
        df_latest = df.loc[mask].copy()
        if len(df_latest) < 5:
            raise ValueError(f"Insufficient data ({len(df_latest)} days)")

        inst_flow = df_latest['institutional_flow']
        ind_flow = df_latest['individual_flow']

        unit, scale_factor = get_unit_and_scale(pd.concat([inst_flow, ind_flow]))

        # Plot 1: Money Flow Trends
        fig1, ax1 = plt.subplots(figsize=(12, 6))
        ax1.plot(df_latest['date'], inst_flow/scale_factor, label='Institutional_Flow')
        ax1.plot(df_latest['date'], ind_flow/scale_factor, label='Individual_Flow')
        ax1.axhline(y=0, color='black', linestyle='--')
        ax1.set_title(f'{current_ticker} Daily Money Flow Trends\n{date_range_label}')
        ax1.set_xlabel('Date')
        ax1.set_ylabel(f'Money Flow ({unit})')
        ax1.legend()
        ax1.grid(alpha=0.3)
        flow_plot = os.path.join(GlobalConfig.full_temp_dir, f"{current_ticker}_flow_trends.png")
        fig1.savefig(flow_plot, bbox_inches='tight')
        plt.close(fig1)
        plot_filenames.append(flow_plot)

        # Plot 2: 5-Day Moving Average
        fig2, ax2 = plt.subplots(figsize=(12, 6))
        if 'inst_5ma' in df_latest.columns and 'ind_5ma' in df_latest.columns:
            ax2.plot(df_latest['date'], df_latest['inst_5ma']/scale_factor, label='Institutional 5-Day MA')
            ax2.plot(df_latest['date'], df_latest['ind_5ma']/scale_factor, label='Individual 5-Day MA')
        else:
            df_latest['inst_5ma'] = df_latest['institutional_flow'].rolling(window=5, min_periods=1).mean()
            df_latest['ind_5ma'] = df_latest['individual_flow'].rolling(window=5, min_periods=1).mean()
            ax2.plot(df_latest['date'], df_latest['inst_5ma']/scale_factor, label='Institutional 5-Day MA')
            ax2.plot(df_latest['date'], df_latest['ind_5ma']/scale_factor, label='Individual 5-Day MA')
        ax2.axhline(y=0, color='black', linestyle='--')
        ax2.set_title(f'{current_ticker} 5-Day Moving Average\n{date_range_label}')
        ax2.set_xlabel('Date')
        ax2.set_ylabel(f'Money Flow ({unit})')
        ax2.legend()
        ax2.grid(alpha=0.3)
        ma_plot = os.path.join(GlobalConfig.full_temp_dir, f"{current_ticker}_moving_averages.png")
        fig2.savefig(ma_plot, bbox_inches='tight')
        plt.close(fig2)
        plot_filenames.append(ma_plot)

        # Pre-calculate all period data for stacked bar charts
        period_data_cache = {}
        all_periods = set([20, 60, 'AP', 5, 10])

        for col in ['extra_large', 'large', 'medium', 'small']:
            if col not in df.columns:
                df[col] = 0

        for period in all_periods:
            if isinstance(period, int) and len(df) >= period:
                df_period = df.tail(period)
                period_label = f"Latest {period} Days"
                period_data_cache[period] = (period_label, df_period)
            elif period == 'AP':
                period_label = f"Analysis Period\n{date_range_label}"
                period_data_cache['AP'] = (period_label, df_latest)

        # Generate stacked bar charts
        period_list1 = [20, 60, 'AP']
        grouped_stacked_plot1 = create_grouped_stacked_bars(
            current_ticker=current_ticker,
            date_range_label=date_range_label,
            period_data_cache=period_data_cache,
            data_periods=period_list1
        )

        if grouped_stacked_plot1:
            plot_filenames.append(grouped_stacked_plot1)

        period_list2 = [5, 10, 20]
        grouped_stacked_plot2 = create_grouped_stacked_bars(
            current_ticker=current_ticker,
            date_range_label=date_range_label,
            period_data_cache=period_data_cache,
            data_periods=period_list2
        )

        if grouped_stacked_plot2:
            plot_filenames.append(grouped_stacked_plot2)

        # Generate analysis report
        analysis_text = [
            f"Analysis Report: {current_ticker}",
            f"Period: {date_range_label}",
            f"Days: {len(df_latest)}"
        ]

        if len(df_latest) > 2:
            try:
                t_stat, p_value = stats.ttest_ind(inst_flow, ind_flow)
                analysis_text.append(f"\nT-stat: {t_stat:.2f}, P-value: {p_value:.4f}")
                analysis_text.append("Significant difference" if p_value < 0.05 else "No significant difference")
            except:
                pass

        analysis_text.append("\nFlow Metrics:")
        if 'idr' in df_latest.columns:
            analysis_text.append(f"IDR: {df_latest['idr'].mean():.2f}")
        if 'fbi' in df_latest.columns:
            analysis_text.append(f"FBI: {df_latest['fbi'].mean():.2f}")
        if 'change_percent' in df_latest.columns:
            analysis_text.append(f"Avg Daily Change: {df_latest['change_percent'].mean():.2f}%")

        # PDF generation
        for attempt in range(3):
            try:
                pdf = FPDF()
                pdf.set_title(f"{current_ticker} Money Flow Analysis")

                pdf.add_page()
                pdf.set_font('Arial', 'B', 16)
                pdf.cell(0, 10, f"{current_ticker} Money Flow Analysis", ln=True, align='C')
                pdf.ln(10)

                pdf.set_font('Arial', '', 12)
                for line in "\n".join(analysis_text).split('\n'):
                    pdf.cell(0, 6, line, ln=True)

                for plot in plot_filenames:
                    if os.path.exists(plot):
                        pdf.add_page()
                        pdf.set_font('Arial', 'B', 14)
                        pdf.cell(0, 10, os.path.basename(plot).replace('_', ' ').replace('.png', ''), ln=True)
                        pdf.image(plot, x=10, w=180)
                        pdf.ln(8)

                pdf.output(pdf_report)

                if os.path.exists(pdf_report) and os.path.getsize(pdf_report) > 1024:
                    pdf_generated = True
                    for plot in plot_filenames:
                        try:
                            os.remove(plot)
                        except:
                            pass
                    break
            except PermissionError:
                time.sleep(1)
            except Exception as e:
                logger.error(f"PDF error: {str(e)}")

        if not pdf_generated:
            raise Exception("PDF generation failed")

        return True

    except Exception as e:
        logger.error(f"Analysis failed: {str(e)}")
        import traceback
        traceback.print_exc()
        return False

def display_money_flow_indicators_from_db(ticker, pdf_report, start_date=None, end_date=None, db_path=None):
    """
    从数据库读取数据，显示资金流指标 - 使用小写列名
    修复版本：解决了PDF合并时的 'endstream' marker 错误
    """
    try:
        df = get_stock_analysis_data(ticker, start_date, end_date, db_path)
        
        if df is None or df.empty:
            logger.error(f"无法获取股票 {ticker} 的数据")
            return

        df.sort_values('date', ascending=True, inplace=True)
        df['date'] = pd.to_datetime(df['date'])

        date_range_label = ""

        if len(df) > 0:
            first_day = df['date'].min()
            last_day = df['date'].max()
        else:
            logger.warning("数据框为空，无法确定日期范围")
            return

        if start_date:
            start_date = pd.to_datetime(start_date)
            if start_date < first_day:
                logger.warning(f"开始日期 {start_date.strftime('%Y-%m-%d')} 早于数据第一天 {first_day.strftime('%Y-%m-%d')}，已自动修正为第一天")
                start_date = first_day
        else:
            start_date = first_day

        if end_date:
            end_date = pd.to_datetime(end_date)        
            if end_date > last_day:
                logger.warning(f"结束日期 {end_date.strftime('%Y-%m-%d')} 超过数据最后日期 {last_day.strftime('%Y-%m-%d')}，已自动修正为最后日期")
                end_date = last_day
        else:
            end_date = last_day

        date_range_label = f"{start_date.strftime('%Y-%m-%d')} 至 {end_date.strftime('%Y-%m-%d')}"

        if len(df) < 3:
            logger.warning("数据不足3天，无法计算连续资金流指标")
            return

        # 确保必要的列存在
        required_cols = ['institutional_net_ratio', 'individual_net_ratio', 
                         'turnover_rate_75th_percentile_500d', 'idr',
                         'change_percent', 'turnover_rate', 'institutional_flow',
                         'individual_flow', 'extra_large', 'large', 'medium', 'small',
                         'volume_ratio5']
        
        for col in required_cols:
            if col not in df.columns:
                if col == 'institutional_net_ratio' and 'idr' in df.columns:
                    df[col] = df['idr'] * 100
                elif col == 'individual_net_ratio' and 'idr' in df.columns:
                    df[col] = (1 - df['idr']) * 100
                else:
                    df[col] = 0
                    logger.warning(f"列 {col} 不存在，已创建默认值")

        # 过滤日期范围
        mask = (df['date'] >= start_date) & (df['date'] <= end_date)
        df_filtered = df.loc[mask].copy()
        
        if len(df_filtered) < 3:
            logger.warning("过滤后的数据不足3天")
            return

        last_10_days = df_filtered.tail(10).copy()

        # 初始化计数变量
        retail_red_count = 0
        retail_red_last_date = None
        retail_green_count = 0
        retail_green_last_date = None
        institution_red_count = 0
        institution_red_last_date = None
        institution_green_count = 0
        institution_green_last_date = None
        inflow_count = 0
        inflow_last_date = None
        outflow_count = 0
        outflow_last_date = None
        multi_inflow_count = 0
        multi_inflow_last_date = None
        multi_outflow_count = 0
        multi_outflow_last_date = None
        heavy_buying_count = 0
        heavy_buying_last_date = None
        selling_count = 0
        selling_last_date = None
        extreme_buying_count = 0
        extreme_buying_last_date = None
        top_divergence_count = 0
        top_divergence_last_date = None
        bottom_divergence_count = 0
        bottom_divergence_last_date = None
        risk_count = 0
        risk_last_date = None

        # 1. 散户主导的阳线
        if all(col in last_10_days.columns for col in ['change_percent', 'turnover_rate', 
                                                       'turnover_rate_75th_percentile_500d', 
                                                       'individual_net_ratio', 'idr']):
            retail_red_condition = (
                (last_10_days['change_percent'] > 0) &
                (last_10_days['turnover_rate'] > last_10_days['turnover_rate_75th_percentile_500d']) &
                (last_10_days['individual_net_ratio'] > 0) &
                (last_10_days['idr']*100 < 20)
            )

            retail_red_occurrences = last_10_days[retail_red_condition]
            retail_red_count = len(retail_red_occurrences)
            retail_red_last_date = retail_red_occurrences['date'].max() if retail_red_count > 0 else None

        # 2. 散户主导的阴线
        if all(col in last_10_days.columns for col in ['change_percent', 'turnover_rate', 
                                                       'turnover_rate_75th_percentile_500d', 
                                                       'individual_net_ratio', 'idr']):
            retail_green_condition = (
                (last_10_days['change_percent'] < 0) &
                (last_10_days['turnover_rate'] > last_10_days['turnover_rate_75th_percentile_500d']) &
                (last_10_days['individual_net_ratio'] < 0) &
                (last_10_days['idr']*100 < 20)
            )

            retail_green_occurrences = last_10_days[retail_green_condition]
            retail_green_count = len(retail_green_occurrences)
            retail_green_last_date = retail_green_occurrences['date'].max() if retail_green_count > 0 else None

        # 3. 机构主导的阳线
        if all(col in last_10_days.columns for col in ['change_percent', 'turnover_rate', 
                                                       'turnover_rate_75th_percentile_500d', 
                                                       'institutional_net_ratio', 'idr']):
            institution_red_condition = (
                (last_10_days['change_percent'] > 0) &
                (last_10_days['turnover_rate'] > last_10_days['turnover_rate_75th_percentile_500d']) &
                (last_10_days['institutional_net_ratio'] > 0) &
                (last_10_days['idr']*100 > 80)
            )

            institution_red_occurrences = last_10_days[institution_red_condition]
            institution_red_count = len(institution_red_occurrences)
            institution_red_last_date = institution_red_occurrences['date'].max() if institution_red_count > 0 else None

        # 4. 机构主导的阴线
        if all(col in last_10_days.columns for col in ['change_percent', 'turnover_rate', 
                                                       'turnover_rate_75th_percentile_500d', 
                                                       'institutional_net_ratio', 'idr']):
            institution_green_condition = (
                (last_10_days['change_percent'] < 0) &
                (last_10_days['turnover_rate'] > last_10_days['turnover_rate_75th_percentile_500d']) &
                (last_10_days['institutional_net_ratio'] < 0) &
                (last_10_days['idr']*100 > 80)
            )

            institution_green_occurrences = last_10_days[institution_green_condition]
            institution_green_count = len(institution_green_occurrences)
            institution_green_last_date = institution_green_occurrences['date'].max() if institution_green_count > 0 else None

        # 5. 主力资金连续3日净流入
        if 'institutional_flow' in last_10_days.columns:
            for i in range(len(last_10_days) - 2):
                window = last_10_days.iloc[i:i+3]
                if all(window['institutional_flow'] > 0):
                    inflow_count += 1
                    inflow_last_date = window.iloc[-1]['date']

        # 6. 主力资金连续3日净流出
        if 'institutional_flow' in last_10_days.columns:
            for i in range(len(last_10_days) - 2):
                window = last_10_days.iloc[i:i+3]
                if all(window['institutional_flow'] < 0):
                    outflow_count += 1
                    outflow_last_date = window.iloc[-1]['date']

        # 7. 多种资金净流入
        required_cols_multi = ['extra_large', 'large', 'medium', 'small']
        if all(col in last_10_days.columns for col in required_cols_multi):
            for i in range(len(last_10_days)):
                day = last_10_days.iloc[i]
                if (day['extra_large'] > 0 and 
                    day['large'] > 0 and 
                    day['medium'] > 0 and 
                    day['small'] > 0):
                    multi_inflow_count += 1
                    multi_inflow_last_date = day['date']

        # 8. 多种资金净流出
        if all(col in last_10_days.columns for col in required_cols_multi):
            for i in range(len(last_10_days)):
                day = last_10_days.iloc[i]
                if (day['extra_large'] < 0 and 
                    day['large'] < 0 and 
                    day['medium'] < 0 and 
                    day['small'] < 0):
                    multi_outflow_count += 1
                    multi_outflow_last_date = day['date']

        # 9. 机构大量吸筹
        if all(col in last_10_days.columns for col in ['institutional_net_ratio', 'volume_ratio5']):
            heavy_buying_condition = (last_10_days['institutional_net_ratio'] > 8) & (last_10_days['volume_ratio5'] > 1.5)
            heavy_buying_occurrences = last_10_days[heavy_buying_condition]
            heavy_buying_count = len(heavy_buying_occurrences)
            if heavy_buying_count > 0:
                heavy_buying_last_date = heavy_buying_occurrences['date'].max()

        # 10. 机构抛售
        if 'institutional_net_ratio' in last_10_days.columns:
            selling_condition = last_10_days['institutional_net_ratio'] < -10
            selling_occurrences = last_10_days[selling_condition]
            selling_count = len(selling_occurrences)
            selling_last_date = selling_occurrences['date'].max() if selling_count > 0 else None

        # 11. 机构极端吸筹
        if 'institutional_net_ratio' in last_10_days.columns:
            extreme_buying_condition = last_10_days['institutional_net_ratio'] > 15
            extreme_buying_occurrences = last_10_days[extreme_buying_condition]
            extreme_buying_count = len(extreme_buying_occurrences)
            extreme_buying_last_date = extreme_buying_occurrences['date'].max() if extreme_buying_count > 0 else None

        # 12. 股价与机构资金流顶背离
        if all(col in last_10_days.columns for col in ['change_percent', 'institutional_flow']):
            for i in range(len(last_10_days) - 2):
                window = last_10_days.iloc[i:i+3]
                if all(window['change_percent'] > 0) and all(window['institutional_flow'] < 0):
                    top_divergence_count += 1
                    top_divergence_last_date = window.iloc[-1]['date']

        # 13. 股价与机构资金流底背离
        if all(col in last_10_days.columns for col in ['change_percent', 'institutional_flow']):
            for i in range(len(last_10_days) - 2):
                window = last_10_days.iloc[i:i+3]
                if all(window['change_percent'] < 0) and all(window['institutional_flow'] > 0):
                    bottom_divergence_count += 1
                    bottom_divergence_last_date = window.iloc[-1]['date']

        # 14. 机构抛售，游资买进（回撤风险高）
        required_cols_risk = ['medium_ratio', 'institutional_net_ratio', 'turnover_rate']
        if all(col in last_10_days.columns for col in required_cols_risk):
            for i in range(len(last_10_days)):
                day = last_10_days.iloc[i]
                if (day['medium_ratio'] > 15 and 
                    day['institutional_net_ratio'] < -3 and 
                    day['turnover_rate'] > 10):
                    risk_count += 1
                    risk_last_date = day['date']

        # 字体设置
        font_path = None
        font_files = [
            'simhei.ttf', 'msyh.ttc', 'msyhbd.ttc', 'simsun.ttc', 'simsunb.ttf'
        ]

        def find_font_file(font_list):
            system_fonts_dirs = [
                'C:/Windows/Fonts', '/usr/share/fonts', '/Library/Fonts', '~/Library/Fonts'
            ]
            for dir in system_fonts_dirs:
                expanded_dir = os.path.expanduser(dir)
                if os.path.exists(expanded_dir):
                    for font in font_list:
                        path = os.path.join(expanded_dir, font)
                        if os.path.exists(path):
                            return path

            local_fonts_dir = 'fonts'
            if os.path.exists(local_fonts_dir):
                for font in font_list:
                    path = os.path.join(local_fonts_dir, font)
                    if os.path.exists(path):
                        return path

            for font in font_list:
                if os.path.exists(font):
                    return font

            return None

        font_path = find_font_file(font_files)

        # 创建PDF
        class IndicatorPDF(FPDF):
            def header(self):
                if font_path:
                    self.set_font('chinese', 'B', 16)
                else:
                    self.set_font('Arial', 'B', 16)
                self.cell(0, 10, f"{ticker} Money Flow Indicators", 0, 1, "C")
                self.ln(5)
                if font_path:
                    self.set_font('chinese', '', 12)
                else:
                    self.set_font('Arial', '', 12)
                self.cell(0, 10, f"Date Range: {date_range_label}", 0, 1, "C")
                self.ln(10)

            def footer(self):
                self.set_y(-15)
                if font_path:
                    self.set_font('chinese', 'I', 8)
                else:
                    self.set_font('Arial', 'I', 8)
                self.cell(0, 10, f"Page {self.page_no()}", 0, 0, "C")

        indicator_pdf = IndicatorPDF()

        # 添加字体
        if font_path:
            try:
                import warnings
                from fpdf import ttfonts
                with warnings.catch_warnings():
                    warnings.filterwarnings("ignore", category=UserWarning, module='fpdf.ttfonts')
                    indicator_pdf.add_font('chinese', '', font_path, uni=True)
                    indicator_pdf.add_font('chinese', 'B', font_path, uni=True)
                    indicator_pdf.add_font('chinese', 'I', font_path, uni=True)
                logger.info(f"字体加载成功: {font_path}")
            except Exception as e:
                logger.warning(f"字体加载失败: {e}")

        indicator_pdf.add_page()

        # 表格列设置
        columns = ["Indicator", "Status", "Details"]
        col_widths = [70, 20, 100]

        # 表头
        if font_path:
            indicator_pdf.set_font('chinese', 'B', 11)
        else:
            indicator_pdf.set_font('Arial', 'B', 11)
            
        for i, col in enumerate(columns):
            indicator_pdf.cell(col_widths[i], 10, col, border=1, align="C")
        indicator_pdf.ln()

        # 准备指标数据
        indicators_data = [
            ("3-Day Institutional Inflow", inflow_count, inflow_last_date, "positive"),
            ("3-Day Institutional Outflow", outflow_count, outflow_last_date, "negative"),
            ("Multi-Class Net Inflow", multi_inflow_count, multi_inflow_last_date, "positive"),
            ("Multi-Class Net Outflow", multi_outflow_count, multi_outflow_last_date, "negative"),
            ("Institutional Heavy Buying", heavy_buying_count, heavy_buying_last_date, "positive"),
            ("Institutional Selling", selling_count, selling_last_date, "negative"),
            ("Institutional Extreme Buying", extreme_buying_count, extreme_buying_last_date, "positive"),
            ("Top Divergence", top_divergence_count, top_divergence_last_date, "negative"),
            ("Bottom Divergence", bottom_divergence_count, bottom_divergence_last_date, "positive"),
            ("Institutional Selling, Hot Money Buying", risk_count, risk_last_date, "negative"),
            ("Retail Dominated Red Candlestick", retail_red_count, retail_red_last_date, "negative"),
            ("Retail Dominated Green Candlestick", retail_green_count, retail_green_last_date, "positive"),
            ("Institution Dominated Red Candlestick", institution_red_count, institution_red_last_date, "positive"),
            ("Institution Dominated Green Candlestick", institution_green_count, institution_green_last_date, "negative"),
        ]

        # 添加数据行
        row_count = 0
        for label, count, last_date, indicator_type in indicators_data:
            if count > 0:
                if font_path:
                    indicator_pdf.set_font('chinese', '', 10)
                else:
                    indicator_pdf.set_font('Arial', '', 10)
                    
                indicator_pdf.cell(col_widths[0], 10, label, border=1)
                
                # 根据指标类型设置颜色
                if indicator_type == "positive":
                    indicator_pdf.set_fill_color(200, 255, 200)  # 绿色
                    status = "True"
                else:
                    indicator_pdf.set_fill_color(255, 200, 200)  # 红色
                    status = "True"
                    
                indicator_pdf.cell(col_widths[1], 10, status, border=1, align="C", fill=True)
                indicator_pdf.set_fill_color(255, 255, 255)
                
                last_date_str = last_date.strftime('%Y-%m-%d') if last_date else "N/A"
                details = f"Count: {count}, Last: {last_date_str}"
                indicator_pdf.cell(col_widths[2], 10, details, border=1)
                indicator_pdf.ln()
                row_count += 1

        # 如果没有指标触发，显示提示信息
        if row_count == 0:
            if font_path:
                indicator_pdf.set_font('chinese', '', 12)
            else:
                indicator_pdf.set_font('Arial', '', 12)
            indicator_pdf.cell(0, 10, "No significant money flow indicators detected in the analysis period.", 0, 1, "C")
            indicator_pdf.ln(5)

        # 添加统计摘要
        indicator_pdf.ln(5)
        if font_path:
            indicator_pdf.set_font('chinese', 'B', 12)
        else:
            indicator_pdf.set_font('Arial', 'B', 12)
        indicator_pdf.cell(0, 10, "Summary Statistics", 0, 1, "L")
        
        if font_path:
            indicator_pdf.set_font('chinese', '', 10)
        else:
            indicator_pdf.set_font('Arial', '', 10)
            
        summary_data = [
            f"Total Days Analyzed: {len(last_10_days)}",
            f"Total Indicators Triggered: {row_count}",
            f"Date Range: {date_range_label}",
            f"Stock Code: {ticker}"
        ]
        
        for line in summary_data:
            indicator_pdf.cell(0, 6, line, 0, 1, "L")

        # ========== 修复：安全地保存并合并PDF ==========
        # 创建临时目录
        temp_dir = os.path.join(GlobalConfig.full_temp_dir, 'temp_pdf')
        os.makedirs(temp_dir, exist_ok=True)
        temp_file = os.path.join(temp_dir, f"{ticker}_money_flow_indicators_temp.pdf")
        
        # 保存临时PDF
        try:
            indicator_pdf.output(temp_file, 'F')
            logger.info(f"临时PDF已生成: {temp_file}")
            
            # 验证临时PDF文件是否有效
            if not os.path.exists(temp_file) or os.path.getsize(temp_file) < 1024:
                logger.error(f"临时PDF文件无效或太小: {temp_file}")
                # 尝试直接保存为最终文件
                fallback_file = os.path.join(GlobalConfig.full_report_dir, f'{ticker}_money_flow_indicators.pdf')
                indicator_pdf.output(fallback_file, 'F')
                logger.info(f"指标PDF已保存为独立文件: {fallback_file}")
                return
                
        except Exception as pdf_error:
            logger.error(f"生成临时PDF失败: {pdf_error}")
            # 尝试直接保存为最终文件
            try:
                fallback_file = os.path.join(GlobalConfig.full_report_dir, f'{ticker}_money_flow_indicators.pdf')
                indicator_pdf.output(fallback_file, 'F')
                logger.info(f"指标PDF已保存为独立文件: {fallback_file}")
            except Exception as fallback_error:
                logger.error(f"保存独立指标PDF失败: {fallback_error}")
            return

        # 合并PDF - 使用更安全的方法
        merged_successfully = False
        try:
            output = PdfWriter()
            
            # 读取现有PDF
            if os.path.exists(pdf_report) and os.path.getsize(pdf_report) > 1024:
                with open(pdf_report, "rb") as f:
                    existing_pdf = PdfReader(f)
                    for page in existing_pdf.pages:
                        output.add_page(page)
                
                # 读取临时PDF并合并
                with open(temp_file, "rb") as f:
                    indicator_pdf_page = PdfReader(f)
                    # 复制所有页面
                    for page in indicator_pdf_page.pages:
                        output.add_page(page)
                
                # 写入最终PDF
                with open(pdf_report, "wb") as f:
                    output.write(f)
                
                merged_successfully = True
                logger.info(f"资金流指标状态已添加到报告: {pdf_report}")
            else:
                logger.warning(f"主PDF不存在或无效，将指标保存为独立文件")
                
        except Exception as merge_error:
            logger.error(f"合并PDF失败: {merge_error}")
            # 不抛出异常，继续执行
            
        # 如果合并失败，保存独立文件
        if not merged_successfully:
            try:
                fallback_file = os.path.join(GlobalConfig.full_report_dir, f'{ticker}_money_flow_indicators.pdf')
                # 使用临时PDF文件复制
                with open(temp_file, "rb") as src:
                    with open(fallback_file, "wb") as dst:
                        dst.write(src.read())
                logger.info(f"指标PDF已保存为独立文件: {fallback_file}")
            except Exception as fallback_error:
                logger.error(f"保存独立指标PDF失败: {fallback_error}")
                # 最后尝试使用indicator_pdf直接输出
                try:
                    fallback_file2 = os.path.join(GlobalConfig.full_report_dir, f'{ticker}_money_flow_indicators_final.pdf')
                    indicator_pdf.output(fallback_file2, 'F')
                    logger.info(f"指标PDF已保存为独立文件(备用): {fallback_file2}")
                except Exception as final_error:
                    logger.error(f"所有保存尝试都失败: {final_error}")
        
        # 清理临时文件
        try:
            if os.path.exists(temp_file):
                os.remove(temp_file)
        except OSError as e:
            logger.warning(f"清理临时文件失败: {e}")

    except Exception as e:
        logger.error(f"添加资金流指标状态时出错: {e}")
        import traceback
        traceback.print_exc()
        
# === 修复的多进程分析函数 ===
def analyze_single_ticker_fixed(ticker_data):
    """
    修复多进程问题的单股票分析函数
    """
    ticker, config_data, start_date, end_date, proj_dir, db_path = ticker_data
    current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    try:
        logger.info(f"{current_time} 开始处理股票: {ticker}")

        GlobalConfig.update_paths(config_data, proj_dir)

        os.makedirs(GlobalConfig.full_data_dir, exist_ok=True)
        os.makedirs(GlobalConfig.full_report_dir, exist_ok=True)
        os.makedirs(GlobalConfig.full_temp_dir, exist_ok=True)

        pdf_report = os.path.join(GlobalConfig.full_report_dir, f'{ticker}_daily_moneyflow_analysis.pdf')

        success = analyze_investor_dominance_from_db(ticker, pdf_report, start_date, end_date, db_path)
        if not success:
            return False, ticker, "投资者主导地位分析失败"

        try:
            display_money_flow_indicators_from_db(ticker, pdf_report, start_date, end_date, db_path)
            logger.info(f"{ticker} 资金流指标状态更新完成")
        except Exception as e:
            logger.warning(f"{ticker} 资金流指标更新失败: {str(e)}")

        logger.info(f"{ticker} 分析完成")
        return True, ticker, "成功"

    except Exception as e:
        logger.error(f"{ticker} 处理失败: {str(e)}")
        return False, ticker, str(e)

# === 简化的多进程版本 ===
def main_simple_multiprocess():
    """
    简化的多进程版本，从 SQLite 数据库读取数据
    """
    try:
        config_path = os.path.join(project_dir, 'Config', 'stock_data_analysis.par')
        CONFIG = load_config(config_path, project_dir)

        GlobalConfig.update_paths(CONFIG, project_dir)

        os.makedirs(GlobalConfig.full_data_dir, exist_ok=True)
        os.makedirs(GlobalConfig.full_report_dir, exist_ok=True)
        os.makedirs(GlobalConfig.full_temp_dir, exist_ok=True)
        
        # ========== 修改：配置日志文件（覆盖模式）==========
        # 确保日志目录存在
        if GlobalConfig.full_log_dir is None:
            GlobalConfig.full_log_dir = os.path.join(project_dir, 'Logs')
        os.makedirs(GlobalConfig.full_log_dir, exist_ok=True)
        
        # 定义日志文件路径
        log_file_path = os.path.join(GlobalConfig.full_log_dir, 'Daily_TA4C_Analyze_MoneyFlow_mproc.log')
        
        # 配置日志处理器 - 同时输出到文件和控制台（覆盖模式）
        log_handlers = [
            logging.FileHandler(log_file_path, mode='w', encoding='utf-8'),  # 添加 mode='w'
            logging.StreamHandler(sys.stdout)
        ]
        
        # 重新配置根日志器
        logging.root.handlers = []
        logging.basicConfig(
            level=logging.INFO,
            format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
            datefmt='%Y-%m-%d %H:%M:%S',  # ← 添加这一行，精确到秒
            handlers=log_handlers
        )
        
        # 更新全局logger
        global logger
        logger = logging.getLogger(__name__)
        
        logger.info(f"日志文件已配置（覆盖模式）: {log_file_path}")
        # ========== 日志配置结束 ==========

        tickers = CONFIG.get('tickers', [])
        start_date, end_date = get_validated_dates(CONFIG, project_dir)

        if not tickers:
            logger.error("没有配置股票代码")
            return

        db_path = os.path.join(project_dir, 'SQLiteDB', 'HK_Stock.db')
        if not os.path.exists(db_path):
            logger.error(f"数据库文件不存在: {db_path}")
            return

        logger.info(f"开始处理 {len(tickers)} 只股票，使用简化多进程模式")
        logger.info(f"日期范围: {start_date.strftime('%Y-%m-%d')} 到 {end_date.strftime('%Y-%m-%d')}")
        logger.info(f"数据库: {db_path}")

        task_data = [(ticker, CONFIG, start_date, end_date, project_dir, db_path) for ticker in tickers]

        successful_tickers = []
        failed_tickers = []

        start_time = time.time()

        max_workers = min(ProcessConfig.MAX_WORKERS, 2)

        logger.info(f"使用 {max_workers} 个进程...")
        
        with concurrent.futures.ProcessPoolExecutor(max_workers=max_workers) as executor:
            future_to_ticker = {
                executor.submit(analyze_single_ticker_fixed, task): task[0] 
                for task in task_data
            }

            completed = 0
            total = len(future_to_ticker)

            for future in concurrent.futures.as_completed(future_to_ticker):
                ticker = future_to_ticker[future]
                completed += 1
                try:
                    success, ticker, message = future.result(timeout=ProcessConfig.TIMEOUT)
                    if success:
                        successful_tickers.append(ticker)
                        logger.info(f"✅ [{completed}/{total}] {ticker} 处理完成")
                    else:
                        failed_tickers.append((ticker, message))
                        logger.error(f"❌ [{completed}/{total}] {ticker} 处理失败: {message}")
                except concurrent.futures.TimeoutError:
                    failed_tickers.append((ticker, "处理超时"))
                    logger.error(f"❌ [{completed}/{total}] {ticker} 处理超时")
                except Exception as e:
                    failed_tickers.append((ticker, str(e)))
                    logger.error(f"❌ [{completed}/{total}] {ticker} 处理异常: {str(e)}")

        end_time = time.time()
        total_time = end_time - start_time

        logger.info("="*50)
        logger.info("处理完成统计:")
        logger.info(f"成功: {len(successful_tickers)} 只股票")
        logger.info(f"失败: {len(failed_tickers)} 只股票")
        logger.info(f"总用时: {total_time:.2f} 秒")
        if len(tickers) > 0:
            logger.info(f"平均每只股票: {total_time/len(tickers):.2f} 秒")

        if failed_tickers:
            logger.info("失败的股票:")
            for ticker, reason in failed_tickers:
                logger.info(f"  {ticker}: {reason}")

        logger.info(f"✅ 全部资金流技术分析报告处理完成！")
        
        # 同时输出到控制台（保持原有打印）
        print("\n" + "="*50)
        print("处理完成统计:")
        print(f"成功: {len(successful_tickers)} 只股票")
        print(f"失败: {len(failed_tickers)} 只股票")
        print(f"总用时: {total_time:.2f} 秒")
        if len(tickers) > 0:
            print(f"平均每只股票: {total_time/len(tickers):.2f} 秒")
        if failed_tickers:
            print("\n失败的股票:")
            for ticker, reason in failed_tickers:
                print(f"  {ticker}: {reason}")
        print(f"\n✅ 全部资金流技术分析报告处理完成！")
        print(f"📝 日志已保存至: {log_file_path}")

    except Exception as e:
        logger.error(f"主程序错误: {str(e)}")
        logger.error(traceback.format_exc())
        print(f"❌ 主程序错误: {str(e)}")
        traceback.print_exc()

# === 单进程版本（备用）===
def main_single_process():
    """
    单进程版本，从 SQLite 数据库读取数据
    """
    try:
        config_path = os.path.join(project_dir, 'Config', 'stock_data_analysis.par')
        CONFIG = load_config(config_path, project_dir)

        GlobalConfig.update_paths(CONFIG, project_dir)

        os.makedirs(GlobalConfig.full_data_dir, exist_ok=True)
        os.makedirs(GlobalConfig.full_report_dir, exist_ok=True)
        os.makedirs(GlobalConfig.full_temp_dir, exist_ok=True)
        
        # ========== 修改：配置日志文件（覆盖模式）==========
        if GlobalConfig.full_log_dir is None:
            GlobalConfig.full_log_dir = os.path.join(project_dir, 'Logs')
        os.makedirs(GlobalConfig.full_log_dir, exist_ok=True)
        
        log_file_path = os.path.join(GlobalConfig.full_log_dir, 'Daily_TA4C_Analyze_MoneyFlow_mproc.log')
        
        log_handlers = [
            logging.FileHandler(log_file_path, mode='w', encoding='utf-8'),  # 添加 mode='w'
            logging.StreamHandler(sys.stdout)
        ]
        
        logging.root.handlers = []
        logging.basicConfig(
            level=logging.INFO,
            format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
            datefmt='%Y-%m-%d %H:%M:%S',  # ← 添加这一行，精确到秒
            handlers=log_handlers
        )
        
        global logger
        logger = logging.getLogger(__name__)
        
        logger.info(f"日志文件已配置（覆盖模式）: {log_file_path}")
        # ========== 日志配置结束 ==========

        tickers = CONFIG.get('tickers', [])
        start_date, end_date = get_validated_dates(CONFIG, project_dir)

        if not tickers:
            logger.error("没有配置股票代码")
            return

        db_path = os.path.join(project_dir, 'SQLiteDB', 'HK_Stock.db')
        if not os.path.exists(db_path):
            logger.error(f"数据库文件不存在: {db_path}")
            return

        logger.info(f"开始单进程处理 {len(tickers)} 只股票")
        logger.info(f"日期范围: {start_date.strftime('%Y-%m-%d')} 到 {end_date.strftime('%Y-%m-%d')}")
        logger.info(f"数据库: {db_path}")

        successful_tickers = []
        failed_tickers = []

        start_time = time.time()

        for i, ticker in enumerate(tickers):
            current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            logger.info(f"[{i+1}/{len(tickers)}] {current_time} 处理股票: {ticker}")
            print(f"[{i+1}/{len(tickers)}] {current_time} 处理股票: {ticker}")

            try:
                pdf_report = os.path.join(GlobalConfig.full_report_dir, f'{ticker}_daily_moneyflow_analysis.pdf')

                success = analyze_investor_dominance_from_db(ticker, pdf_report, start_date, end_date, db_path)
                if not success:
                    failed_tickers.append((ticker, "投资者主导地位分析失败"))
                    logger.error(f"❌ {ticker} 投资者主导地位分析失败")
                    print(f"❌ {ticker} 投资者主导地位分析失败")
                    continue

                try:
                    display_money_flow_indicators_from_db(ticker, pdf_report, start_date, end_date, db_path)
                    logger.info(f"{ticker} 资金流指标状态更新完成")
                except Exception as e:
                    logger.warning(f"{ticker} 资金流指标更新失败: {str(e)}")

                successful_tickers.append(ticker)
                logger.info(f"✅ [{i+1}/{len(tickers)}] {ticker} 处理完成")
                print(f"✅ [{i+1}/{len(tickers)}] {ticker} 处理完成")

            except Exception as e:
                logger.error(f"{ticker} 处理失败: {str(e)}")
                failed_tickers.append((ticker, str(e)))
                print(f"❌ [{i+1}/{len(tickers)}] {ticker} 处理失败: {str(e)}")

        end_time = time.time()
        total_time = end_time - start_time

        logger.info("="*50)
        logger.info("处理完成统计:")
        logger.info(f"成功: {len(successful_tickers)} 只股票")
        logger.info(f"失败: {len(failed_tickers)} 只股票")
        logger.info(f"总用时: {total_time:.2f} 秒")
        if len(tickers) > 0:
            logger.info(f"平均每只股票: {total_time/len(tickers):.2f} 秒")

        if failed_tickers:
            logger.info("失败的股票:")
            for ticker, reason in failed_tickers:
                logger.info(f"  {ticker}: {reason}")

        logger.info(f"✅ 全部资金流技术分析报告处理完成！")
        
        print("\n" + "="*50)
        print("处理完成统计:")
        print(f"成功: {len(successful_tickers)} 只股票")
        print(f"失败: {len(failed_tickers)} 只股票")
        print(f"总用时: {total_time:.2f} 秒")
        if len(tickers) > 0:
            print(f"平均每只股票: {total_time/len(tickers):.2f} 秒")
        if failed_tickers:
            print("\n失败的股票:")
            for ticker, reason in failed_tickers:
                print(f"  {ticker}: {reason}")
        print(f"\n✅ 全部资金流技术分析报告处理完成！")
        print(f"📝 日志已保存至: {log_file_path}")

    except Exception as e:
        logger.error(f"主程序错误: {str(e)}")
        logger.error(traceback.format_exc())
        print(f"❌ 主程序错误: {str(e)}")
        traceback.print_exc()

# === 主函数选择 ===
def main():
    """主函数 - 提供多种运行模式"""
    if sys.platform.startswith('win'):
        mp.freeze_support()

    main_simple_multiprocess()

if __name__ == "__main__":
    main()